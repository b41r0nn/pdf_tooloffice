"""
conversion_excel.py — PDF → Excel

Motores (todos libres), elegidos con un banco de pruebas de facturas, reportes
de inventario, extractos y escaneos:
  - PyMuPDF find_tables (líneas): tablas con bordes dibujados; conserva celdas
    de varias líneas sin partirlas.
  - Camelot 2 (MIT), modo 'hybrid': tablas sin bordes (columnas alineadas) y
    tablas cuyas líneas son imagen (escaneos).
  - Tesseract (integrado en MuPDF) para páginas escaneadas: se les agrega una
    capa de texto OCR antes de extraer.

El Excel resultante tiene:
  Resumen  qué se encontró, con qué método y avisos
  Tabla N  una hoja por tabla (las que siguen en varias páginas se unen y se
           quitan los encabezados repetidos), como Tabla de Excel con filtros
  Datos    pares "Etiqueta: valor" fuera de las tablas (NIT, fecha, total...)
  Texto    todo el texto fuera de las tablas, línea por línea
Los valores se convierten a número, fecha o porcentaje por columna, detectando
si el documento usa coma o punto decimal. Los códigos (ceros a la izquierda,
columnas Código/NIT/Lote/Referencia...) se dejan como texto.
"""

import io
import logging
import os
import re
import tempfile
import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import datetime

import openpyxl
import pymupdf as fitz  # PyMuPDF
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as TablaExcel
from openpyxl.worksheet.table import TableStyleInfo

try:
    import camelot
except Exception:  # pragma: no cover - dependencia opcional
    camelot = None

try:  # vienen con Camelot; se usan para limpiar la imagen antes del OCR
    import cv2
    import numpy as np
except Exception:  # pragma: no cover - dependencia opcional
    cv2 = np = None

log = logging.getLogger(__name__)

MAX_PAGINAS_OCR = 40
DPI_OCR = 300
MAX_CELDA_EXCEL = 32767
MAX_VALOR_DATO = 60

METODO_BORDES = 'Tabla con bordes (PyMuPDF)'
METODO_CAMELOT = 'Tabla sin bordes (Camelot)'
METODO_TEXTO = 'Alineación de texto (PyMuPDF)'


# ── Modelo ────────────────────────────────────────────────────────

@dataclass
class Tabla:
    paginas: list
    filas: list                       # filas de datos (listas de str)
    encabezado: list = None           # nombres de columna o None
    metodo: str = ''
    zonas: list = field(default_factory=list)   # [(pagina, bbox PyMuPDF)] por cada página que ocupa


@dataclass
class Contenido:
    paginas: int = 0
    tablas: list = field(default_factory=list)
    texto: list = field(default_factory=list)     # (pagina, linea, y) — y para ordenar en la página
    datos: list = field(default_factory=list)     # (pagina, campo, valor_str)
    avisos: list = field(default_factory=list)
    paginas_ocr: list = field(default_factory=list)


# ── Limpieza de texto ─────────────────────────────────────────────

def limpiar(texto):
    """Normaliza una celda: une líneas partidas por el ancho de columna y quita artefactos."""
    if texto is None:
        return ''
    t = str(texto).replace('\xa0', ' ').replace('\u2010', '-').replace('\u2011', '-').replace('\r', '\n')
    t = re.sub(r'-\n(?=[a-záéíóúñ])', '', t)        # palabra cortada con guion al final de línea
    t = re.sub(r'\s*\n\s*', ' ', t)
    t = re.sub(r'(^|\s)[|¦]+(?=\s|$)', ' ', t)       # bordes leídos como '|' por el OCR
    t = re.sub(r'\s+', ' ', t).strip()
    return t


def _sin_acentos(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn').lower()


# ── Números y fechas ──────────────────────────────────────────────

_RE_MONEDA = re.compile(r'(?i)(US\$|COP\$?|USD|EUR|\$|€|£)')
_RE_NUCLEO = re.compile(r'^\d[\d.,\s]*$')
_MESES = {
    'ene': 1, 'enero': 1, 'jan': 1, 'feb': 2, 'febrero': 2, 'mar': 3, 'marzo': 3,
    'abr': 4, 'abril': 4, 'apr': 4, 'may': 5, 'mayo': 5, 'jun': 6, 'junio': 6,
    'jul': 7, 'julio': 7, 'ago': 8, 'agosto': 8, 'aug': 8, 'sep': 9, 'sept': 9,
    'septiembre': 9, 'set': 9, 'oct': 10, 'octubre': 10, 'nov': 11, 'noviembre': 11,
    'dic': 12, 'diciembre': 12, 'dec': 12,
}
# Primera palabra del encabezado que indica una columna de códigos (se deja como texto)
_PALABRAS_CODIGO = {'codigo', 'cod', 'ref', 'referencia', 'nit', 'cedula', 'cc', 'documento', 'doc',
                    'lote', 'cuenta', 'telefono', 'tel', 'celular', 'no', 'nro', 'numero', 'num', 'serie',
                    'consecutivo', 'cufe', 'pedido', 'orden', 'item', 'ean', 'barras', 'cum', 'invima',
                    'factura', 'id', 'placa', 'serial'}


def _votar_decimal(textos):
    """',' o '.' según cómo aparecen los separadores en los números; None si no hay evidencia."""
    coma = punto = 0
    for t in textos:
        nucleo = _RE_MONEDA.sub('', t).replace('%', '').strip(' ()+-')
        if not _RE_NUCLEO.match(nucleo):
            continue
        nucleo = nucleo.replace(' ', '')
        if '.' in nucleo and ',' in nucleo:
            if nucleo.rfind(',') > nucleo.rfind('.'):
                coma += 1
            else:
                punto += 1
            continue
        for sep in ('.', ','):
            if sep not in nucleo:
                continue
            partes = nucleo.split(sep)
            if len(partes) > 2:          # 1.250.000 -> separador de miles
                if sep == '.':
                    coma += 1
                else:
                    punto += 1
            elif len(partes[1]) != 3:   # 450,5 o 8750,25 -> separador decimal
                if sep == ',':
                    coma += 1
                else:
                    punto += 1
    if coma == punto == 0:
        return None
    return ',' if coma >= punto else '.'


def a_numero(texto, decimal=','):
    """
    Convierte texto a (valor, formato_excel) o None si no es un número.
    Acepta moneda, signo, paréntesis negativos y porcentaje. No convierte
    códigos con ceros a la izquierda ni cadenas de más de 15 dígitos.
    """
    t = texto.strip()
    if not t or not any(ch.isdigit() for ch in t):
        return None
    negativo = (t.startswith('(') and t.endswith(')')) or t.startswith('-') or t.endswith('-')
    porcentaje = t.endswith('%') or t.startswith('%')
    moneda = bool(_RE_MONEDA.search(t))
    nucleo = _RE_MONEDA.sub('', t).replace('%', '').strip().strip('()').strip().strip('+-').strip()
    if not _RE_NUCLEO.match(nucleo):
        return None
    nucleo = nucleo.replace(' ', '')
    miles = '.' if decimal == ',' else ','
    if nucleo.count(decimal) > 1:
        return None
    entero, _, fraccion = nucleo.partition(decimal)
    if miles in entero:
        grupos = entero.split(miles)
        if not (1 <= len(grupos[0]) <= 3 and all(len(g) == 3 for g in grupos[1:])):
            return None
        entero = ''.join(grupos)
    if miles in fraccion or not entero.isdigit() or (fraccion and not fraccion.isdigit()):
        return None
    if len(entero) > 1 and entero.startswith('0') and not fraccion:
        return None                     # código: 000123
    if len(entero) > 15:
        return None                     # Excel pierde precisión: cédulas/CUFE como texto
    valor = float(f'{entero}.{fraccion}') if fraccion else int(entero)
    if negativo:
        valor = -valor
    if porcentaje:
        decimales = len(fraccion)
        return valor / 100, '0%' if decimales == 0 else '0.' + '0' * decimales + '%'
    if fraccion:
        return valor, '#,##0.' + '0' * len(fraccion)
    if moneda or abs(valor) >= 1000 or miles in nucleo:
        return valor, '#,##0'
    return valor, 'General'


_RE_FECHA_NUM = re.compile(r'^(\d{1,4})[/.-](\d{1,2})[/.-](\d{1,4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$')
_RE_FECHA_TXT = re.compile(r'^(\d{1,2})(?:\s+de)?[\s/.-]+([a-záéíóú]{3,10})\.?(?:\s+de)?[\s/.-]+(\d{2,4})$', re.I)


def a_fecha(texto, dia_primero=True):
    """Convierte texto a datetime (día primero, como en Colombia) o None."""
    t = texto.strip()
    m = _RE_FECHA_NUM.match(t)
    try:
        if m:
            a, b, c = (int(x) for x in m.group(1, 2, 3))
            if len(m.group(1)) == 4:
                anio, mes, dia = a, b, c
            elif len(m.group(3)) in (2, 4):
                anio = c + (2000 if c < 70 else 1900) if c < 100 else c
                dia, mes = (a, b) if dia_primero else (b, a)
            else:
                return None
            hora = tuple(int(x) for x in m.group(4, 5) if x) + ((int(m.group(6)),) if m.group(6) else ())
            return datetime(anio, mes, dia, *hora)
        m = _RE_FECHA_TXT.match(t)
        if m:
            mes = _MESES.get(_sin_acentos(m.group(2)).rstrip('.'))
            if not mes:
                return None
            anio = int(m.group(3))
            anio = anio + 2000 if anio < 100 else anio
            return datetime(anio, mes, int(m.group(1)))
    except ValueError:
        return None
    return None


def _orden_fechas(textos):
    """True si las fechas van día/mes (lo usual en Colombia); False si mes/día."""
    for t in textos:
        m = _RE_FECHA_NUM.match(t.strip())
        if m and len(m.group(1)) <= 2:
            if int(m.group(1)) > 12:
                return True
            if int(m.group(2)) > 12:
                return False
    return True


def _es_columna_codigo(nombre):
    palabras = re.findall(r'[a-z0-9]+', _sin_acentos(nombre or ''))
    return bool(palabras) and palabras[0] in _PALABRAS_CODIGO


def tipar_columna(valores, nombre='', decimal=','):
    """
    Decide el tipo de una columna y retorna [(valor, formato)] del mismo largo.
    Solo convierte si al menos el 80 % de las celdas no vacías encajan, para no
    mezclar tipos dentro de una columna.
    """
    no_vacias = [v for v in valores if v]
    salida = [(v, None) for v in valores]
    if not no_vacias or _es_columna_codigo(nombre):
        return salida
    dia_primero = _orden_fechas(no_vacias)
    fechas = [a_fecha(v, dia_primero) for v in valores]
    if sum(1 for f in fechas if f) >= 0.8 * len(no_vacias):
        con_hora = any(f and (f.hour or f.minute) for f in fechas)
        formato = 'dd/mm/yyyy hh:mm' if con_hora else 'dd/mm/yyyy'
        return [((f, formato) if f else (v, None)) for v, f in zip(valores, fechas)]
    numeros = [a_numero(v, decimal) for v in valores]
    if sum(1 for n in numeros if n) >= 0.8 * len(no_vacias):
        formato = _formato_columna([n for n in numeros if n])
        return [((n[0], formato) if n else (v, None)) for v, n in zip(valores, numeros)]
    return salida


def _formato_columna(numeros):
    """Un solo formato para toda la columna (el más detallado de sus celdas)."""
    formatos = [f for _, f in numeros]
    porcentajes = [f for f in formatos if f.endswith('%')]
    if len(porcentajes) == len(formatos):
        return max(porcentajes, key=len)
    decimales = max((len(f.split('.')[1]) for f in formatos if f.startswith('#,##0.')), default=0)
    if decimales:
        return '#,##0.' + '0' * decimales
    if any(f == '#,##0' for f in formatos):
        return '#,##0'
    return 'General'


def tipar_valor(texto, decimal=','):
    """Tipo de un valor suelto (hoja Datos)."""
    f = a_fecha(texto)
    if f:
        return f, 'dd/mm/yyyy'
    n = a_numero(texto, decimal)
    return n if n else (texto, None)


# ── OCR ───────────────────────────────────────────────────────────

def ruta_tessdata():
    """Carpeta con spa.traineddata (Tesseract) o None si el OCR no está disponible."""
    candidatas = [os.environ.get('TESSDATA_PREFIX'),
                  '/usr/share/tesseract-ocr/5/tessdata', '/usr/share/tesseract-ocr/4.00/tessdata',
                  '/usr/share/tessdata', r'C:\Program Files\Tesseract-OCR\tessdata']
    for d in candidatas:
        if d and os.path.isfile(os.path.join(d, 'spa.traineddata')):
            return d
    return None


def es_escaneada(page):
    """Página que es básicamente una imagen (puede traer un pie de página con texto)."""
    if len(page.get_text().strip()) >= 40:
        return False
    area = sum(abs(fitz.Rect(i['bbox']) & page.rect) for i in page.get_image_info())
    return area >= 0.5 * abs(page.rect)


def _quitar_lineas(pix):
    """
    Copia de la página sin las líneas de tablas, solo para el OCR: Tesseract no
    reconoce el texto encerrado en cajas con poco margen y las trata como figuras.
    """
    img = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)
    gris = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
    binaria = cv2.threshold(gris, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    horizontales = cv2.morphologyEx(binaria, cv2.MORPH_OPEN,
                                    cv2.getStructuringElement(cv2.MORPH_RECT, (max(20, pix.width // 30), 1)))
    verticales = cv2.morphologyEx(binaria, cv2.MORPH_OPEN,
                                  cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(20, pix.height // 40))))
    mascara = cv2.dilate(cv2.bitwise_or(horizontales, verticales), np.ones((3, 3), np.uint8))
    gris[mascara > 0] = 255
    limpio = fitz.Pixmap(fitz.csRGB, pix.width, pix.height, cv2.cvtColor(gris, cv2.COLOR_GRAY2RGB).tobytes(), False)
    limpio.set_dpi(pix.xres, pix.yres)
    return limpio


def _ocr_pagina(page, idioma, tessdata):
    """Documento de 1 página: imagen original + capa de texto OCR invisible."""
    pix = page.get_pixmap(dpi=DPI_OCR, colorspace=fitz.csRGB, alpha=False)
    para_ocr = pix
    if cv2 is not None:
        try:
            para_ocr = _quitar_lineas(pix)
        except Exception:
            log.debug('No se pudieron quitar las líneas antes del OCR', exc_info=True)
    ocr = fitz.open('pdf', para_ocr.pdfocr_tobytes(language=idioma, tessdata=tessdata))
    if para_ocr is not pix:
        imagenes = ocr[0].get_images()
        if imagenes:            # se deja la imagen con líneas: Camelot las usa para ubicar la tabla
            ocr[0].replace_image(imagenes[0][0], pixmap=pix)
    return ocr


def _preparar(doc, contenido):
    """Agrega capa OCR a las páginas escaneadas. Retorna un documento nuevo o None si no hizo falta."""
    escaneadas = [p.number for p in doc if es_escaneada(p)]
    if not escaneadas:
        return None
    tessdata = ruta_tessdata()
    if not tessdata:
        contenido.avisos.append(f'{len(escaneadas)} página(s) escaneada(s) sin texto: el OCR no está '
                                'instalado (tesseract-ocr-spa), no se pudo leer su contenido.')
        return None
    if len(escaneadas) > MAX_PAGINAS_OCR:
        contenido.avisos.append(f'Se aplicó OCR solo a las primeras {MAX_PAGINAS_OCR} páginas escaneadas.')
        escaneadas = escaneadas[:MAX_PAGINAS_OCR]
    idioma = 'spa+eng' if os.path.isfile(os.path.join(tessdata, 'eng.traineddata')) else 'spa'
    nuevo = fitz.open()
    for page in doc:
        if page.number in escaneadas:
            with _ocr_pagina(page, idioma, tessdata) as ocr:
                nuevo.insert_pdf(ocr)
        else:
            nuevo.insert_pdf(doc, from_page=page.number, to_page=page.number)
    contenido.paginas_ocr = [n + 1 for n in escaneadas]
    contenido.avisos.append('Páginas escaneadas leídas con OCR (revisar cifras): '
                            + ', '.join(map(str, contenido.paginas_ocr)))
    return nuevo


# ── Tablas ────────────────────────────────────────────────────────

def _normalizar_filas(filas):
    """Limpia celdas y quita filas y columnas totalmente vacías."""
    filas = [[limpiar(c) for c in f] for f in filas]
    filas = [f for f in filas if any(f)]
    if not filas:
        return []
    ancho = max(len(f) for f in filas)
    filas = [f + [''] * (ancho - len(f)) for f in filas]
    usadas = [i for i in range(ancho) if any(f[i] for f in filas)]
    return [[f[i] for i in usadas] for f in filas]


def _es_encabezado(fila):
    llenas = [c for c in fila if c]
    if len(llenas) < max(2, len(fila) // 2):
        return False
    def es_dato(c):
        if re.fullmatch(r'(19|20)\d{2}', c):      # años como encabezado: 2025 | 2024
            return False
        return bool(a_numero(c) or a_numero(c, '.') or a_fecha(c))
    if any(es_dato(c) for c in llenas):
        return False
    return any(re.search(r'[A-Za-zÁÉÍÓÚÑáéíóúñ]', c) for c in llenas)


def _tabla_valida(filas):
    if len(filas) < 2 or len(filas[0]) < 2:
        return False
    una_celda = sum(1 for f in filas if sum(1 for c in f if c) <= 1)
    if una_celda > 0.6 * len(filas):
        return False                    # texto corrido detectado como tabla
    largo = [len(c) for f in filas for c in f if c]
    return sum(largo) / len(largo) < 80


def _crear_tabla(filas, pagina, metodo, bbox, contenido):
    filas = _normalizar_filas(filas)
    # Títulos de una sola celda antes del encabezado van a las hojas Texto/Datos
    while filas and sum(1 for c in filas[0] if c) <= 1 and len(filas) > 2:
        _agregar_linea(contenido, pagina, [next((c for c in filas[0] if c), '')], bbox[1] - 1 if bbox else 0)
        filas.pop(0)
    if not _tabla_valida(filas):
        return None
    encabezado = filas.pop(0) if _es_encabezado(filas[0]) else None
    return Tabla(paginas=[pagina], filas=filas, encabezado=encabezado, metodo=metodo,
                 zonas=[(pagina, bbox)] if bbox else [])


def _bbox_camelot(tabla_camelot, alto_pagina):
    caja = getattr(tabla_camelot, '_bbox', None)   # origen abajo-izquierda (PDF)
    if not caja:
        return None
    x1, y1, x2, y2 = caja
    return (x1, alto_pagina - y2, x2, alto_pagina - y1)


def _extraer_tablas(doc, ruta_trabajo, contenido):
    tablas = []
    sin_tablas = []
    for page in doc:
        encontradas = []
        try:
            for t in page.find_tables(strategy='lines').tables:
                tabla = _crear_tabla(t.extract(), page.number + 1, METODO_BORDES, tuple(t.bbox), contenido)
                if tabla:
                    encontradas.append(tabla)
        except Exception:
            log.debug('find_tables falló en la página %d', page.number + 1, exc_info=True)
        if encontradas:
            tablas += encontradas
        else:
            sin_tablas.append(page.number)

    if sin_tablas and camelot is not None:
        paginas = ','.join(str(n + 1) for n in sin_tablas)
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                resultado = camelot.read_pdf(ruta_trabajo, pages=paginas, flavor='hybrid')
            for t in resultado:
                n = int(t.page)
                tabla = _crear_tabla(t.df.values.tolist(), n, METODO_CAMELOT,
                                     _bbox_camelot(t, doc[n - 1].rect.height), contenido)
                if tabla:
                    tablas.append(tabla)
        except Exception:
            log.warning('Camelot no pudo procesar las páginas %s', paginas, exc_info=True)
    elif sin_tablas:
        contenido.avisos.append('Camelot no está instalado: las tablas sin bordes se detectan con un '
                                'método menos preciso.')
        for n in sin_tablas:
            page = doc[n]
            try:
                for t in page.find_tables(strategy='text').tables:
                    tabla = _crear_tabla(t.extract(), n + 1, METODO_TEXTO, tuple(t.bbox), contenido)
                    if tabla:
                        tablas.append(tabla)
            except Exception:
                log.debug('find_tables(text) falló en la página %d', n + 1, exc_info=True)

    tablas = _quitar_superpuestas(tablas)
    tablas.sort(key=lambda t: (t.paginas[0], t.zonas[0][1][1] if t.zonas else 0))
    return tablas


def _quitar_superpuestas(tablas):
    """Camelot puede devolver la misma tabla dos veces (una parcial): se deja la de más celdas."""
    resultado = []
    for t in sorted(tablas, key=lambda t: -sum(len(f) for f in t.filas)):
        caja = fitz.Rect(t.zonas[0][1]) if t.zonas else None
        duplicada = caja is not None and any(
            o.paginas[0] == t.paginas[0] and o.zonas
            and abs(caja & fitz.Rect(o.zonas[0][1])) > 0.5 * min(abs(caja), abs(fitz.Rect(o.zonas[0][1])))
            for o in resultado)
        if not duplicada:
            resultado.append(t)
    return resultado


def _unir_continuaciones(tablas):
    """Une tablas que continúan en la página siguiente (mismas columnas) y quita encabezados repetidos."""
    unidas = []
    for t in tablas:
        previa = unidas[-1] if unidas else None
        continua = (
            previa is not None
            and t.paginas[0] == previa.paginas[-1] + 1
            and len((t.encabezado or t.filas[0])) == len((previa.encabezado or previa.filas[0]))
            and (t.encabezado is None or previa.encabezado is None or t.encabezado == previa.encabezado)
        )
        if continua:
            filas = t.filas
            if t.encabezado is None and previa.encabezado and filas and filas[0] == previa.encabezado:
                filas = filas[1:]
            previa.filas += filas
            previa.paginas.append(t.paginas[0])
            previa.zonas += t.zonas
        else:
            unidas.append(t)
    return unidas


# ── Texto y datos fuera de tablas ─────────────────────────────────

# Varios espacios seguidos de otra etiqueta ('Firma: ___   Fecha: 01/10/2026') separan dos datos
_RE_CORTE_ETIQUETA = re.compile(r'\s{2,}(?=[A-ZÁÉÍÓÚÑ][\w .áéíóúñ]{0,30}:\s)')
_RE_PAR = re.compile(r'^\s*([A-Za-zÁÉÍÓÚÑáéíóúñ][\w .#º°/()áéíóúñÁÉÍÓÚÑ-]{1,45}?)\s*:\s*(\S.*)$')


def _segmentos_linea(line):
    """Divide una línea en segmentos separados por espacios grandes (columnas de un encabezado)."""
    segmentos = []
    actual = ''
    fin_prev = None
    for span in line['spans']:
        texto = span['text']
        if not texto.strip():
            continue
        x0 = span['bbox'][0]
        if fin_prev is not None:
            hueco = x0 - fin_prev
            if hueco > max(14, span['size'] * 1.6):
                segmentos.append(actual)
                actual = ''
            elif hueco > span['size'] * 0.12 and actual and not actual.endswith(' ') and not texto.startswith(' '):
                actual += ' '           # el OCR guarda cada palabra en su propio span, sin espacios
        actual += texto
        fin_prev = span['bbox'][2]
    if actual.strip():
        segmentos.append(actual)
    resultado = []
    for seg in segmentos:
        for parte in _RE_CORTE_ETIQUETA.split(seg):     # 'Firma: ____   Fecha: 01/10/2026'
            parte = limpiar(parte)
            if parte:
                resultado.append(parte)
    return resultado


def _texto_y_datos(doc, tablas, contenido):
    por_pagina = {}
    for t in tablas:
        for pagina, bbox in t.zonas:
            por_pagina.setdefault(pagina, []).append(fitz.Rect(bbox))
    for page in doc:
        n = page.number + 1
        zonas = [r + (-2, -2, 2, 2) for r in por_pagina.get(n, [])]
        for block in page.get_text('dict', sort=True)['blocks']:
            if block.get('type') != 0:
                continue
            for line in block['lines']:
                r = fitz.Rect(line['bbox'])
                centro = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
                if any(centro in z for z in zonas):
                    continue
                segmentos = _segmentos_linea(line)
                if segmentos:
                    _agregar_linea(contenido, n, segmentos, r.y0)


def _agregar_linea(contenido, pagina, segmentos, y=0):
    """Agrega una línea a Texto y sus pares 'Etiqueta: valor' a Datos."""
    contenido.texto.append((pagina, '   '.join(segmentos), y))
    for seg in segmentos:
        for parte in re.split(r'\s[·•|]\s', seg):              # 'Titular: X · Periodo: Y'
            m = _RE_PAR.match(parte)
            # Valores largos son la primera línea de un párrafo ("PRIMERA. OBJETO: El proveedor..."), no datos
            if (m and len(m.group(2)) <= MAX_VALOR_DATO and not m.group(2).startswith('//')
                    and not re.fullmatch(r'[_.\-–\s]+', m.group(2))):   # '____' = campo para firmar
                contenido.datos.append((pagina, m.group(1).strip(), m.group(2).strip()))


# ── Extracción completa ───────────────────────────────────────────

def extraer_contenido(ruta_pdf):
    """Extrae tablas, texto y datos de un PDF (sin escribir el Excel)."""
    contenido = Contenido()
    with fitz.open(ruta_pdf) as original:
        if original.needs_pass:
            raise ValueError('El PDF está protegido con contraseña')
        contenido.paginas = original.page_count
        con_ocr = _preparar(original, contenido)
        doc = con_ocr or original
        ruta_trabajo = ruta_pdf
        tmp = None
        try:
            if con_ocr is not None:
                fd, tmp = tempfile.mkstemp(suffix='.pdf')
                os.close(fd)
                con_ocr.save(tmp)
                ruta_trabajo = tmp
            tablas = _unir_continuaciones(_extraer_tablas(doc, ruta_trabajo, contenido))
            contenido.tablas = tablas
            _texto_y_datos(doc, tablas, contenido)
            contenido.texto.sort(key=lambda x: (x[0], x[2]))
            contenido.datos.sort(key=lambda x: x[0])
        finally:
            if con_ocr is not None:
                con_ocr.close()
            if tmp:
                try:
                    os.remove(tmp)
                except OSError:
                    pass
    if not contenido.tablas:
        contenido.avisos.append('No se detectaron tablas: el contenido está en las hojas Datos y Texto.')
    return contenido


# ── Excel ─────────────────────────────────────────────────────────

_RELLENO_ENC = PatternFill(start_color='1B4FD8', end_color='1B4FD8', fill_type='solid')
_FUENTE_ENC = Font(bold=True, color='FFFFFF')


def _texto_excel(valor):
    if isinstance(valor, str):
        valor = ILLEGAL_CHARACTERS_RE.sub('', valor)[:MAX_CELDA_EXCEL]
    return valor


def _escribir(ws, fila, col, valor, formato=None):
    cell = ws.cell(row=fila, column=col, value=_texto_excel(valor))
    if isinstance(valor, str):
        cell.data_type = 's'            # texto que empieza por '=' no es fórmula
        if len(valor) > 60:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    if formato and formato != 'General':
        cell.number_format = formato
    return cell


def _nombres_unicos(encabezado, ancho):
    nombres = []
    for i in range(ancho):
        base = (encabezado[i] if encabezado and i < len(encabezado) else '') or f'Columna {i + 1}'
        base = _texto_excel(base)[:250]
        nombre, k = base, 2
        while nombre in nombres:
            nombre = f'{base} ({k})'
            k += 1
        nombres.append(nombre)
    return nombres


def _ajustar_anchos(ws, columnas, filas_muestra):
    for i in range(columnas):
        largo = max((len(str(f[i])) for f in filas_muestra if i < len(f) and f[i] not in (None, '')), default=8)
        ws.column_dimensions[get_column_letter(i + 1)].width = min(60, max(9, largo + 2))


def _hoja_tabla(wb, indice, tabla, decimal_doc):
    ws = wb.create_sheet(f'Tabla {indice}')
    ancho = max(len(tabla.encabezado or []), *(len(f) for f in tabla.filas))
    filas = [f + [''] * (ancho - len(f)) for f in tabla.filas]
    nombres = _nombres_unicos(tabla.encabezado, ancho)
    decimal = _votar_decimal([c for f in filas for c in f]) or decimal_doc

    for c, nombre in enumerate(nombres, start=1):
        cell = _escribir(ws, 1, c, nombre)
        cell.font, cell.fill = _FUENTE_ENC, _RELLENO_ENC
    columnas = [tipar_columna([f[c] for f in filas], nombres[c], decimal) for c in range(ancho)]
    for r in range(len(filas)):
        for c in range(ancho):
            valor, formato = columnas[c][r]
            if valor != '':
                _escribir(ws, r + 2, c + 1, valor, formato)

    ref = f'A1:{get_column_letter(ancho)}{len(filas) + 1}'
    tabla_xl = TablaExcel(displayName=f'Tabla{indice}', ref=ref)
    tabla_xl.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabla_xl)
    ws.freeze_panes = 'A2'
    _ajustar_anchos(ws, ancho, [nombres] + filas[:200])
    return ws


def _hoja_simple(wb, titulo, encabezado, filas, nombre_tabla):
    ws = wb.create_sheet(titulo)
    for c, nombre in enumerate(encabezado, start=1):
        cell = _escribir(ws, 1, c, nombre)
        cell.font, cell.fill = _FUENTE_ENC, _RELLENO_ENC
    for r, fila in enumerate(filas, start=2):
        for c, (valor, formato) in enumerate(fila, start=1):
            if valor not in (None, ''):
                _escribir(ws, r, c, valor, formato)
    if filas:
        t = TablaExcel(displayName=nombre_tabla, ref=f'A1:{get_column_letter(len(encabezado))}{len(filas) + 1}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = 'A2'
    _ajustar_anchos(ws, len(encabezado), [encabezado] + [[v for v, _ in f] for f in filas[:200]])
    return ws


def escribir_excel(contenido, nombre_archivo=''):
    wb = openpyxl.Workbook()
    resumen = wb.active
    resumen.title = 'Resumen'

    todas = [c for t in contenido.tablas for f in t.filas for c in f] + [v for _, _, v in contenido.datos]
    decimal_doc = _votar_decimal(todas) or ','

    for i, tabla in enumerate(contenido.tablas, start=1):
        _hoja_tabla(wb, i, tabla, decimal_doc)
    if contenido.datos:
        _hoja_simple(wb, 'Datos', ['Página', 'Campo', 'Valor'],
                     [[(p, None), (campo, None), tipar_valor(valor, decimal_doc)] for p, campo, valor in contenido.datos],
                     'Datos')
    if contenido.texto:
        _hoja_simple(wb, 'Texto', ['Página', 'Texto'], [[(p, None), (t, None)] for p, t, _ in contenido.texto], 'Texto')

    # Resumen
    filas = [('Archivo', nombre_archivo or '—'), ('Páginas', contenido.paginas),
             ('Tablas encontradas', len(contenido.tablas)),
             ('Separador decimal detectado', 'coma (1.234,56)' if decimal_doc == ',' else 'punto (1,234.56)')]
    if contenido.paginas_ocr:
        filas.append(('Páginas leídas con OCR', ', '.join(map(str, contenido.paginas_ocr))))
    r = 1
    _escribir(resumen, r, 1, 'Conversión PDF → Excel').font = Font(bold=True, size=13)
    r += 2
    for clave, valor in filas:
        _escribir(resumen, r, 1, clave).font = Font(bold=True)
        _escribir(resumen, r, 2, valor)
        r += 1
    if contenido.tablas:
        r += 1
        for c, nombre in enumerate(['Hoja', 'Páginas', 'Filas', 'Columnas', 'Método'], start=1):
            cell = _escribir(resumen, r, c, nombre)
            cell.font, cell.fill = _FUENTE_ENC, _RELLENO_ENC
        for i, t in enumerate(contenido.tablas, start=1):
            r += 1
            paginas = f'{t.paginas[0]}–{t.paginas[-1]}' if len(t.paginas) > 1 else str(t.paginas[0])
            ancho = max(len(t.encabezado or []), *(len(f) for f in t.filas))
            for c, valor in enumerate([f'Tabla {i}', paginas, len(t.filas), ancho, t.metodo], start=1):
                _escribir(resumen, r, c, valor)
    if contenido.avisos:
        r += 2
        _escribir(resumen, r, 1, 'Avisos').font = Font(bold=True, color='B45309')
        for aviso in contenido.avisos:
            r += 1
            _escribir(resumen, r, 1, aviso)
    resumen.column_dimensions['A'].width = 30
    resumen.column_dimensions['B'].width = 40
    for col in 'CDE':
        resumen.column_dimensions[col].width = 14
    resumen.column_dimensions['E'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_a_xlsx(ruta_pdf, nombre_archivo=''):
    return escribir_excel(extraer_contenido(ruta_pdf), nombre_archivo)
