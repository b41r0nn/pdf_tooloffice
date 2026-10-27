import io
import re
import zipfile

import pymupdf as fitz
import openpyxl
import pytest
from docx import Document
from PIL import Image

from app import services
from app.database import registrar_log

from .conftest import ADMIN, pdf_bytes, textos


def _archivo(datos, nombre):
    return (io.BytesIO(datos), nombre)


@pytest.fixture(autouse=True)
def sin_libreoffice(monkeypatch):
    # Resultados deterministas: se prueba el renderizado básico.
    monkeypatch.setattr(services, 'ruta_libreoffice', lambda: None)


# ── Generales ─────────────────────────────────────────────────────

def test_health_y_cabeceras(client):
    r = client.get('/health')
    assert r.status_code == 200 and r.json['status'] == 'ok'
    assert r.headers['X-Content-Type-Options'] == 'nosniff'
    assert r.headers['X-Frame-Options'] == 'SAMEORIGIN'


def test_index_usa_rutas_estaticas_y_version(client):
    html = client.get('/').get_data(as_text=True)
    assert '/static/js/editor.js' in html and '/static/pdfjs/web/viewer.html' in html
    from app import __version__
    assert f'v{__version__}' in html
    # CSS/JS con ?v=<fecha> para no quedar en caché tras un despliegue; fuentes sin él (preload = URL del CSS)
    assert re.search(r'/static/css/style\.css\?v=\d+', html)
    assert '/static/fonts/bricolage-grotesque.woff2"' in html


def test_ajustes_del_visor_van_en_hoja_aparte(client):
    # La CSP de viewer.html (style-src 'self') bloquea <style> en línea: editor.js enlaza visor.css
    html = client.get('/').get_data(as_text=True)
    m = re.search(r'data-css="(/static/css/visor\.css\?v=\d+)"', html)
    assert m
    r = client.get(m.group(1))
    assert r.status_code == 200 and r.mimetype == 'text/css'
    assert b'#mainContainer' in r.data
    r.close()
    visor = client.get('/static/pdfjs/web/viewer.html').get_data(as_text=True)
    assert "style-src 'self'" in visor


def test_pdfjs_se_sirve_como_modulo_js(client):
    r = client.get('/static/pdfjs/build/pdf.mjs')
    assert r.status_code == 200
    assert r.mimetype == 'text/javascript'
    r.close()
    assert client.get('/static/js/main.js').mimetype == 'text/javascript'


def test_favicons_y_manifiesto(client):
    html = client.get('/').get_data(as_text=True)
    for ruta, tipo in (('icons/favicon.ico', 'image/x-icon'), ('icons/icon.svg', 'image/svg+xml'),
                       ('icons/apple-touch-icon.png', 'image/png'),
                       ('icons/manifest.webmanifest', 'application/manifest+json'),
                       ('fonts/bricolage-grotesque.woff2', 'font/woff2')):
        if not ruta.startswith('fonts/'):
            assert f'/static/{ruta}' in html
        r = client.get(f'/static/{ruta}')
        assert r.status_code == 200 and r.mimetype == tipo, ruta
        r.close()
    r = client.get('/favicon.ico')
    assert r.status_code == 200 and r.mimetype == 'image/x-icon'
    r.close()
    manifiesto = client.get('/static/icons/manifest.webmanifest').get_json(force=True)
    for icono in manifiesto['icons']:
        r = client.get(f"/static/icons/{icono['src']}")
        assert r.status_code == 200, icono['src']
        r.close()


def test_archivo_demasiado_grande_responde_json(app, client):
    app.config['MAX_CONTENT_LENGTH'] = 1024
    r = client.post('/api/paginas', data={'archivo': _archivo(b'x' * 4096, 'a.pdf')})
    assert r.status_code == 413 and 'límite' in r.json['error']


def test_pdf_falso_rechazado(client):
    r = client.post('/api/paginas', data={'archivo': _archivo(b'no soy pdf', 'falso.pdf')})
    assert r.status_code == 400


# ── Admin ─────────────────────────────────────────────────────────

def test_admin_escapa_html_de_los_logs(app, client):
    registrar_log(app.config['DATABASE_PATH'], 'division', 'x.pdf', 10, 'error',
                  'Modo no soportado: <script>alert(1)</script>')
    assert client.get('/admin/logs').status_code == 401
    assert client.get('/admin/logs', headers={'X-Admin-Password': 'mala'}).status_code == 401
    r = client.get('/admin/logs', headers={'X-Admin-Password': ADMIN})
    html = r.get_data(as_text=True)
    assert r.status_code == 200
    assert '<script>alert(1)</script>' not in html
    assert '&lt;script&gt;' in html


def test_admin_deshabilitado_sin_contrasena(app, client):
    app.config['ADMIN_PASSWORD'] = ''
    assert client.get('/admin/logs', headers={'X-Admin-Password': ''}).status_code == 401


# ── Operaciones ───────────────────────────────────────────────────

def test_union_conserva_orden_y_nombre(client):
    r = client.post('/api/union', data={
        'archivos': [_archivo(pdf_bytes([('Primero', 100)]), 'a.pdf'), _archivo(pdf_bytes([('Segundo', 100)]), 'b.pdf')],
        'nombre_salida': 'Unión final.pdf',
    })
    assert r.status_code == 200
    assert textos(r.data, 0) == ['Primero'] and textos(r.data, 1) == ['Segundo']
    assert "Uni%C3%B3n%20final.pdf" in r.headers['Content-Disposition']


def test_union_rechaza_nivel_invalido(client):
    r = client.post('/api/union', data={
        'archivos': [_archivo(pdf_bytes(), 'a.pdf'), _archivo(pdf_bytes(), 'b.pdf')],
        'comprimir': 'true', 'nivel': '<script>',
    })
    assert r.status_code == 400


def test_division_rango_y_modos(client):
    datos = pdf_bytes(paginas=5)
    r = client.post('/api/division', data={'archivo': _archivo(datos, 'd.pdf'), 'modo': 'rango', 'inicio': '2', 'fin': '4'})
    assert r.status_code == 200
    assert fitz.open('pdf', r.data).page_count == 3
    r = client.post('/api/division', data={'archivo': _archivo(datos, 'd.pdf'), 'modo': 'rango'})
    assert r.status_code == 400
    r = client.post('/api/division', data={'archivo': _archivo(datos, 'd.pdf'), 'modo': 'pares'})
    assert fitz.open('pdf', r.data).page_count == 2
    r = client.post('/api/division', data={'archivo': _archivo(datos, 'd.pdf'), 'modo': 'todas'})
    assert len(zipfile.ZipFile(io.BytesIO(r.data)).namelist()) == 5


def test_division_cada_entrega_un_archivo_por_bloque(client):
    r = client.post('/api/division', data={'archivo': _archivo(pdf_bytes(paginas=10), 'escaneo.pdf'),
                                           'modo': 'cada', 'paginas_por_archivo': '2'})
    assert r.status_code == 200
    assert r.mimetype == 'application/zip'
    zf = zipfile.ZipFile(io.BytesIO(r.data))
    assert zf.namelist() == [f'escaneo_{i:02d}.pdf' for i in range(1, 6)]
    assert all(fitz.open('pdf', zf.read(n)).page_count == 2 for n in zf.namelist())


def test_division_personalizado_respeta_los_grupos_y_su_orden(client):
    r = client.post('/api/division', data={'archivo': _archivo(pdf_bytes(paginas=10), 'tanda.pdf'),
                                           'modo': 'personalizado', 'grupos': '6-7,1-2,3-5'})
    assert r.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(r.data))
    assert [fitz.open('pdf', zf.read(n)).page_count for n in zf.namelist()] == [2, 2, 3]


def test_division_personalizado_deja_fuera_las_paginas_no_listadas(client):
    r = client.post('/api/division', data={'archivo': _archivo(pdf_bytes(paginas=10), 't.pdf'),
                                           'modo': 'personalizado', 'grupos': '1-3'})
    zf = zipfile.ZipFile(io.BytesIO(r.data))
    assert len(zf.namelist()) == 1
    assert fitz.open('pdf', zf.read(zf.namelist()[0])).page_count == 3


@pytest.mark.parametrize('datos', [
    {'modo': 'cada', 'paginas_por_archivo': '0'},
    {'modo': 'cada', 'paginas_por_archivo': '-3'},
    {'modo': 'cada', 'paginas_por_archivo': 'abc'},
    {'modo': 'cada'},
    {'modo': 'personalizado', 'grupos': '1-4,3-6'},
    {'modo': 'personalizado', 'grupos': '1-2,40-50'},
    {'modo': 'personalizado', 'grupos': ''},
    {'modo': 'personalizado'},
    {'modo': 'inventado'},
])
def test_division_rechaza_entradas_invalidas(client, datos):
    r = client.post('/api/division', data={'archivo': _archivo(pdf_bytes(paginas=10), 'd.pdf'), **datos})
    assert r.status_code == 400
    assert 'error' in r.get_json()


def test_division_cada_con_pdf_cifrado_pide_la_contrasena(client):
    cifrado = client.post('/api/cifrado', data={'archivo': _archivo(pdf_bytes(paginas=4), 'c.pdf'),
                                                'contrasena': 'clave', 'confirmar': 'clave'}).data
    r = client.post('/api/division', data={'archivo': _archivo(cifrado, 'c.pdf'),
                                           'modo': 'cada', 'paginas_por_archivo': '2'})
    assert r.status_code == 400
    assert 'contraseña' in r.get_json()['error']


def test_division_nombre_con_ruta_no_escapa_del_zip(client):
    r = client.post('/api/division', data={'archivo': _archivo(pdf_bytes(paginas=4), '../../factura.pdf'),
                                           'modo': 'cada', 'paginas_por_archivo': '2'})
    nombres = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert nombres == ['factura_01.pdf', 'factura_02.pdf']
    assert not any('..' in n or '/' in n or '\\' in n for n in nombres)


def test_cifrado_aes256(client):
    r = client.post('/api/cifrado', data={'archivo': _archivo(pdf_bytes(), 'c.pdf'),
                                          'contrasena': 'clave', 'confirmar': 'clave'})
    assert r.status_code == 200
    doc = fitz.open('pdf', r.data)
    assert doc.needs_pass
    assert doc.authenticate('clave')
    assert 'AES' in doc.metadata['encryption'] and '256' in doc.metadata['encryption']


@pytest.mark.parametrize('nivel', ['baja', 'media', 'alta', 'ultra'])
def test_compresion_conserva_tamano_de_pagina(client, nivel):
    with fitz.open() as doc:
        page = doc.new_page(width=595, height=842)
        img = Image.effect_noise((1600, 1600), 60).convert('RGB')
        buf = io.BytesIO()
        img.save(buf, 'PNG')
        page.insert_image(fitz.Rect(50, 50, 545, 545), stream=buf.getvalue())
        original = doc.tobytes()
    r = client.post('/api/compresion', data={'archivo': _archivo(original, 'grande.pdf'), 'nivel': nivel})
    assert r.status_code == 200
    resultado = fitz.open('pdf', r.data)
    assert resultado[0].rect == fitz.Rect(0, 0, 595, 842)
    assert len(r.data) <= len(original)
    if nivel in ('alta', 'ultra'):
        assert len(r.data) < len(original) * 0.5


def test_conversion_pdf_a_word_con_extension_mayuscula(client):
    r = client.post('/api/conversion', data={'archivo': _archivo(pdf_bytes(), 'REPORTE.PDF'), 'direccion': 'pdf-docx'})
    assert r.status_code == 200
    assert 'Fecha' in '\n'.join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)


def test_conversion_rechaza_extension_que_no_corresponde(client):
    r = client.post('/api/conversion', data={'archivo': _archivo(b'PK', 'datos.xlsx'), 'direccion': 'pdf-docx'})
    assert r.status_code == 400


def test_word_a_pdf_con_caracteres_especiales_y_tablas(client):
    doc = Document()
    doc.add_heading('Informe & resumen <2026>', level=1)
    doc.add_paragraph('Texto con & y < y >')
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = 'Producto'
    tabla.cell(1, 0).text = 'Guantes'
    tabla.cell(1, 1).text = '$ 1.000'
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post('/api/conversion', data={'archivo': _archivo(buf.getvalue(), 'w.docx'), 'direccion': 'docx-pdf'})
    assert r.status_code == 200
    texto = ' '.join(textos(r.data))
    assert 'Informe & resumen <2026>' in texto and 'Guantes' in texto


def test_excel_a_pdf_muestra_valores_y_todas_las_hojas(client):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Item', 'Valor'])
    ws.append(['A', 5])
    ws['B3'] = '=SUM(B2:B2)'
    wb.create_sheet('Segunda').append(['Hoja dos'])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post('/api/conversion', data={'archivo': _archivo(buf.getvalue(), 'x.xlsx'), 'direccion': 'xlsx-pdf'})
    assert r.status_code == 200
    with fitz.open('pdf', r.data) as pdf:
        todo = ' '.join(p.get_text() for p in pdf)
    assert 'Hoja dos' in todo
    assert '=SUM' not in todo


def test_imagenes_a_pdf_transparencia_y_tiff_multipagina(client):
    png = io.BytesIO()
    Image.new('RGBA', (100, 100), (0, 0, 0, 0)).save(png, 'PNG')
    tif = io.BytesIO()
    Image.new('RGB', (80, 120), 'red').save(tif, 'TIFF', save_all=True,
                                             append_images=[Image.new('RGB', (80, 120), 'blue')])
    r = client.post('/api/conversion', data={
        'direccion': 'imagen_a_pdf',
        'archivos': [_archivo(png.getvalue(), 'transparente.png'), _archivo(tif.getvalue(), 'escaneo.tif')],
    })
    assert r.status_code == 200
    doc = fitz.open('pdf', r.data)
    assert doc.page_count == 3
    centro = doc[0].get_pixmap(dpi=20).pixel(80, 110)
    assert min(centro) > 240  # fondo blanco, no negro


def test_pdf_a_excel_neutraliza_formulas(client):
    r = client.post('/api/conversion', data={'archivo': _archivo(pdf_bytes([('=HYPERLINK("http://x")', 100)]), 'f.pdf'),
                                             'direccion': 'pdf-xlsx'})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    celdas = [c for ws in wb.worksheets for fila in ws.iter_rows() for c in fila
              if isinstance(c.value, str) and c.value.startswith('=')]
    assert celdas and all(c.data_type == 's' for c in celdas)


# ── Editor ────────────────────────────────────────────────────────

def _crear_sesion(client, datos=None, nombre='Factura Ñ.pdf', **extra):
    r = client.post('/api/edicion', data={'archivo': _archivo(datos or pdf_bytes(), nombre), **extra})
    assert r.status_code == 201, r.json
    return r.json


def test_editor_flujo_completo(client):
    s = _crear_sesion(client, pdf_bytes([('Fecha: 2026-08-05', 100), ('Linea vecina', 111)]))
    assert s['version'] == 1 and s['paginas'] == 1 and not s['puede_deshacer']

    r = client.get(f"/api/edicion/{s['id']}/paginas/0/campos")
    campo = next(c for c in r.json['campos'] if c['texto'].startswith('Fecha'))

    r = client.post(f"/api/edicion/{s['id']}/reemplazar", json={
        'pagina': 0, 'bbox': campo['bbox'], 'texto_original': campo['texto'], 'texto_nuevo': 'Fecha: 2026-12-31'})
    assert r.status_code == 200 and r.json['version'] == 2

    r = client.get(f"/api/edicion/{s['id']}/pdf")
    assert set(textos(r.data)) == {'Fecha: 2026-12-31', 'Linea vecina'}

    r = client.get(f"/api/edicion/{s['id']}/paginas/0/imagen?dpi=72")
    assert r.status_code == 200 and r.mimetype == 'image/png'
    assert client.get(f"/api/edicion/{s['id']}/paginas/0/imagen?dpi=5000").status_code == 400

    r = client.post(f"/api/edicion/{s['id']}/deshacer")
    assert r.json['version'] == 1
    r = client.get(f"/api/edicion/{s['id']}/descargar")
    assert 'Fecha: 2026-08-05' in textos(r.data)
    assert "Factura%20%C3%91_editado.pdf" in r.headers['Content-Disposition']

    assert client.delete(f"/api/edicion/{s['id']}").status_code == 204
    assert client.get(f"/api/edicion/{s['id']}").status_code == 404


def test_editor_guarda_version_del_visor_y_aplana(client):
    s = _crear_sesion(client)
    with fitz.open('pdf', pdf_bytes()) as doc:
        annot = doc[0].add_freetext_annot(fitz.Rect(72, 200, 300, 230), 'Firmado: Juan', fontsize=12)
        annot.update()
        anotado = doc.tobytes()
    r = client.post(f"/api/edicion/{s['id']}/version", data={'archivo': _archivo(anotado, 'visor.pdf')})
    assert r.status_code == 200 and r.json['version'] == 2

    sin_aplanar = client.get(f"/api/edicion/{s['id']}/descargar").data
    assert len(list(fitz.open('pdf', sin_aplanar)[0].annots())) == 1
    aplanado = client.get(f"/api/edicion/{s['id']}/descargar?aplanar=1").data
    doc = fitz.open('pdf', aplanado)
    assert len(list(doc[0].annots())) == 0
    assert 'Firmado: Juan' in doc[0].get_text()


def test_editor_rechaza_ids_maliciosos(client):
    for ruta in ('/api/edicion/..%2F..%2Fsecreto/pdf', '/api/edicion/' + 'a' * 32 + '/pdf',
                 '/api/edicion/xyz/descargar'):
        assert client.get(ruta).status_code == 404
    assert client.post('/api/edicion/' + '0' * 32 + '/reemplazar', json={}).status_code == 404


def test_editor_valida_parametros(client):
    s = _crear_sesion(client)
    url = f"/api/edicion/{s['id']}/reemplazar"
    assert client.post(url, json={'pagina': 0, 'bbox': [1, 2, 3], 'texto_nuevo': 'x'}).status_code == 400
    assert client.post(url, json={'pagina': '0', 'bbox': [1, 2, 3, 4], 'texto_nuevo': 'x'}).status_code == 400
    assert client.post(url, json={'pagina': 0, 'bbox': [72, 87, 170, 104], 'texto_nuevo': 'x',
                                  'alineacion': 'otra'}).status_code == 400
    assert client.post(f"/api/edicion/{s['id']}/deshacer").status_code == 400


def test_editor_pdf_protegido(client):
    with fitz.open('pdf', pdf_bytes()) as doc:
        cifrado = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='abc', owner_pw='abc')
    r = client.post('/api/edicion', data={'archivo': _archivo(cifrado, 'p.pdf')})
    assert r.status_code == 400
    s = _crear_sesion(client, cifrado, contrasena='abc')
    assert s['desbloqueado'] is True
    assert not fitz.open('pdf', client.get(f"/api/edicion/{s['id']}/pdf").data).needs_pass
