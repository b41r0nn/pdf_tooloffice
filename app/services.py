"""
services.py — Lógica de procesamiento PDF
Librerías: PyMuPDF, pdf2docx, python-docx, openpyxl, reportlab y Pillow.
PDF → Excel está en conversion_excel.py (PyMuPDF + Camelot + OCR).

Conversión Office → PDF: si LibreOffice está instalado (o SOFFICE_PATH apunta
a él) se usa en modo headless, que conserva tablas, estilos y saltos de
página. Si no, se usa un renderizado básico con reportlab.
"""

import io
import logging
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from xml.sax.saxutils import escape

import pymupdf as fitz  # PyMuPDF
import openpyxl
from docx import Document
from docx.table import Table as DocxTable
from pdf2docx import Converter

from .conversion_excel import es_escaneada, extraer_contenido, ruta_tessdata
from .conversion_excel import pdf_a_xlsx  # noqa: F401  (PDF → Excel vive en su propio módulo)
from PIL import Image, ImageOps, ImageSequence
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    import magic
except Exception:  # pragma: no cover - dependencia opcional
    magic = None

log = logging.getLogger(__name__)

NIVELES_COMPRESION = ('baja', 'media', 'alta', 'ultra')
MODOS_DIVISION = ('todas', 'pares', 'impares')
MAX_ARCHIVOS_DIVISION = 500
MAX_TEXTO_GRUPOS = 10_000

# ── Validación ────────────────────────────────────────────────────

def validar_pdf(ruta_pdf):
    """
    Valida que el archivo sea un PDF legible.
    Retorna (es_valido, mensaje_error)
    """
    if magic:
        try:
            mime = magic.from_file(ruta_pdf, mime=True)
            if mime not in ('application/pdf', 'application/x-pdf'):
                return False, f'MIME type inválido: {mime}'
        except Exception:
            pass  # Continuar sin magic si falla

    # La especificación admite basura antes del header dentro del primer KB.
    with open(ruta_pdf, 'rb') as f:
        if b'%PDF-' not in f.read(1024):
            return False, 'Archivo no es un PDF válido (header inválido)'

    try:
        with fitz.open(ruta_pdf) as doc:
            if not doc.is_pdf:
                return False, 'Archivo no es un PDF válido'
            if doc.page_count == 0 and not doc.needs_pass:
                return False, 'PDF no contiene páginas'
    except Exception:
        return False, 'El PDF está dañado o no se puede leer'
    return True, ''


_PARTE_OFFICE = {'docx': 'word/document.xml', 'xlsx': 'xl/workbook.xml'}


def validar_mime_type(ruta_archivo, tipo_esperado):
    """
    Valida el tipo real del archivo.
    tipo_esperado puede ser: 'pdf', 'docx', 'xlsx'
    Retorna (es_valido, mensaje_error)

    DOCX/XLSX se validan por su estructura (paquete ZIP con la parte principal):
    libmagic identifica como application/zip u octet-stream muchos archivos
    válidos generados por programas distintos a Office.
    """
    if tipo_esperado in _PARTE_OFFICE:
        try:
            with zipfile.ZipFile(ruta_archivo) as zf:
                nombres = set(zf.namelist())
        except (zipfile.BadZipFile, OSError):
            return False, f'El archivo no es un .{tipo_esperado} válido'
        if '[Content_Types].xml' not in nombres or _PARTE_OFFICE[tipo_esperado] not in nombres:
            return False, f'El archivo no es un .{tipo_esperado} válido'
        return True, ''

    if tipo_esperado != 'pdf' or not magic:
        return True, ''  # Sin magic, validar_pdf() revisa header y estructura

    try:
        mime = magic.from_file(ruta_archivo, mime=True)
        if mime not in ('application/pdf', 'application/x-pdf'):
            return False, f'MIME type inválido. Esperado: pdf, encontrado: {mime}'
        return True, ''
    except Exception:
        return True, ''  # Si falla magic, aceptar

# ── Utilidades ────────────────────────────────────────────────────

def abrir_pdf(ruta, contrasena=''):
    """Abre un PDF y, si está cifrado, lo desbloquea con la contraseña."""
    doc = fitz.open(ruta)
    if doc.needs_pass and not (contrasena and doc.authenticate(contrasena)):
        doc.close()
        raise ValueError('El archivo PDF está cifrado: ingresa la contraseña correcta para procesarlo')
    return doc


def contar_paginas(ruta_pdf, contrasena=''):
    es_valido, msg_error = validar_pdf(ruta_pdf)
    if not es_valido:
        raise ValueError(msg_error)
    with abrir_pdf(ruta_pdf, contrasena) as doc:
        return doc.page_count


def _guardar(doc, **extra):
    return doc.tobytes(garbage=3, deflate=True, **extra)


# ── Unión ─────────────────────────────────────────────────────────

def unir_pdfs(rutas, contrasena=''):
    for ruta in rutas:
        es_valido, msg_error = validar_pdf(ruta)
        if not es_valido:
            raise ValueError(f'PDF inválido: {msg_error}')

    with fitz.open() as salida:
        for ruta in rutas:
            with abrir_pdf(ruta, contrasena) as src:
                salida.insert_pdf(src)
        return _guardar(salida)


# ── División ──────────────────────────────────────────────────────

def dividir_pdf(ruta, inicio, fin):
    with abrir_pdf(ruta) as src:
        total = src.page_count
        inicio = max(1, inicio)
        fin = min(fin, total)
        if inicio > fin:
            raise ValueError(f'Rango inválido: {inicio}-{fin} (total {total} páginas)')
        with fitz.open() as salida:
            salida.insert_pdf(src, from_page=inicio - 1, to_page=fin - 1)
            return _guardar(salida)


def dividir_pdf_por_modo(ruta, modo):
    if modo not in MODOS_DIVISION:
        raise ValueError('Modo de división no soportado. Opciones: '
                         f'rango, {", ".join(MODOS_DIVISION)}, cada, personalizado')

    with abrir_pdf(ruta) as src:
        total = src.page_count
        if total == 0:
            raise ValueError('El PDF no contiene páginas')

        if modo == 'todas':
            ancho = len(str(total))
            out = io.BytesIO()
            with zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
                for i in range(total):
                    with fitz.open() as pagina:
                        pagina.insert_pdf(src, from_page=i, to_page=i)
                        zf.writestr(f'pagina_{i + 1:0{ancho}d}.pdf', _guardar(pagina))
            return out.getvalue(), 'zip'

        resto = 0 if modo == 'pares' else 1
        indices = [i for i in range(total) if (i + 1) % 2 == resto]
        if not indices:
            raise ValueError('No hay páginas para el modo seleccionado')
        src.select(indices)
        return _guardar(src), 'pdf'


def analizar_grupos(texto, total):
    """
    Convierte '1-2,3-5,6' en [(1, 2), (3, 5), (6, 6)] (páginas desde 1, ambas incluidas,
    en el orden escrito). Se pueden dejar páginas fuera y escribir los grupos en cualquier
    orden, pero una página no puede estar en dos grupos.
    Lanza ValueError con un mensaje pensado para quien lo lee en pantalla.
    """
    if len(texto) > MAX_TEXTO_GRUPOS:
        raise ValueError('El texto de los grupos es demasiado largo.')

    trozos = [t.strip() for t in texto.split(',')]
    trozos = [t for t in trozos if t]
    if not trozos:
        raise ValueError('Escribe al menos un grupo de páginas, por ejemplo 1-2,3-5,6.')

    grupos = []
    for trozo in trozos:
        m = re.fullmatch(r'([0-9]{1,6})(?:-([0-9]{1,6}))?', trozo)
        if not m:
            visto = trozo if len(trozo) <= 40 else trozo[:40] + '…'
            raise ValueError(f'«{visto}» no es un grupo válido; escribe un número o un tramo como 3-5.')
        inicio = int(m.group(1))
        fin = int(m.group(2)) if m.group(2) else inicio
        grupos.append((inicio, fin))

    for inicio, fin in grupos:
        for pagina in (inicio, fin):
            if pagina < 1 or pagina > total:
                raise ValueError(f'La página {pagina} no existe; el documento tiene de la 1 a la {total}.')

    for inicio, fin in grupos:
        if inicio > fin:
            raise ValueError(f'El grupo {inicio}-{fin} está al revés: la página inicial es mayor que la final.')

    if len(grupos) > MAX_ARCHIVOS_DIVISION:
        raise ValueError(f'Hay {len(grupos)} grupos, más de los {MAX_ARCHIVOS_DIVISION} permitidos.')

    usadas = set()
    for inicio, fin in grupos:
        for pagina in range(inicio, fin + 1):
            if pagina in usadas:
                raise ValueError(f'La página {pagina} está en más de un grupo.')
            usadas.add(pagina)

    return grupos


def grupos_por_bloques(total, paginas_por_archivo):
    """Parte `total` páginas en bloques de `paginas_por_archivo`; el último puede quedar más corto."""
    if paginas_por_archivo < 1:
        raise ValueError('Las páginas por archivo deben ser al menos 1.')

    cantidad = -(-total // paginas_por_archivo)  # división hacia arriba
    if cantidad > MAX_ARCHIVOS_DIVISION:
        raise ValueError(
            f'Saldrían {cantidad} archivos, más de los {MAX_ARCHIVOS_DIVISION} permitidos. '
            'Usa «Cada página por separado».'
        )

    return [
        (inicio, min(inicio + paginas_por_archivo - 1, total))
        for inicio in range(1, total + 1, paginas_por_archivo)
    ]


def dividir_por_grupos(ruta, grupos, base='documento'):
    """ZIP con un PDF por grupo (inicio, fin; páginas desde 1, ambas incluidas), en el orden dado."""
    if not grupos:
        raise ValueError('No hay páginas para dividir')

    ancho = max(2, len(str(len(grupos))))
    out = io.BytesIO()
    with abrir_pdf(ruta) as src, zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        for numero, (inicio, fin) in enumerate(grupos, start=1):
            with fitz.open() as salida:
                salida.insert_pdf(src, from_page=inicio - 1, to_page=fin - 1)
                zf.writestr(f'{base}_{numero:0{ancho}d}.pdf', _guardar(salida))
    return out.getvalue()


# ── Conversión ────────────────────────────────────────────────────

def _imagen_a_rgb(img):
    """Convierte a RGB; la transparencia se compone sobre fondo blanco (no negro)."""
    if img.mode in ('RGBA', 'LA') or (img.mode == 'P' and 'transparency' in img.info):
        img = img.convert('RGBA')
        fondo = Image.new('RGB', img.size, (255, 255, 255))
        fondo.paste(img, mask=img.getchannel('A'))
        return fondo
    if img.mode != 'RGB':
        return img.convert('RGB')
    return img


def convertir_imagenes_a_pdf(archivos: list) -> bytes:
    """Cada imagen (o cada página de un TIFF multipágina) va a una página A4 centrada."""
    A4_W, A4_H = fitz.paper_size('a4')

    with fitz.open() as doc:
        for archivo in archivos:
            try:
                img = Image.open(archivo)
                frames = [frame.copy() for frame in ImageSequence.Iterator(img)]
            except Exception as e:
                log.warning('Imagen ilegible %s: %s', os.path.basename(archivo), e)
                continue

            for frame in frames:
                try:
                    frame = _imagen_a_rgb(ImageOps.exif_transpose(frame))
                except Exception as e:
                    log.warning('No se pudo procesar %s: %s', os.path.basename(archivo), e)
                    continue

                # Orientación de la página según la imagen
                target_w, target_h = (A4_H, A4_W) if frame.width > frame.height else (A4_W, A4_H)
                page = doc.new_page(width=target_w, height=target_h)

                # Escala máxima sin deformar, centrada
                ratio = min(target_w / frame.width, target_h / frame.height)
                new_w, new_h = frame.width * ratio, frame.height * ratio
                x_offset = (target_w - new_w) / 2
                y_offset = (target_h - new_h) / 2
                rect = fitz.Rect(x_offset, y_offset, x_offset + new_w, y_offset + new_h)

                buf = io.BytesIO()
                frame.save(buf, format='JPEG', quality=92)
                page.insert_image(rect, stream=buf.getvalue())

        if doc.page_count == 0:
            raise ValueError('Ninguna de las imágenes se pudo procesar')
        return _guardar(doc)


def convertir_pdf_a_jpg(ruta: str, dpi: int = 150):
    """Retorna (bytes, paginas): JPG si es una página, ZIP de JPGs si son varias."""
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    with abrir_pdf(ruta) as doc:
        paginas = doc.page_count
        if paginas == 1:
            return doc[0].get_pixmap(matrix=mat, alpha=False).tobytes('jpeg', jpg_quality=90), 1

        ancho = len(str(paginas))
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            for i, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=mat, alpha=False)
                zf.writestr(f'pagina_{i:0{ancho}d}.jpg', pix.tobytes('jpeg', jpg_quality=90))
        return out.getvalue(), paginas


def _docx_desde_contenido(contenido):
    """Word editable a partir del texto y las tablas extraídos (con OCR)."""
    documento = Document()
    if contenido.paginas_ocr:
        aviso = documento.add_paragraph()
        aviso.add_run('Documento reconstruido con OCR: revise nombres y cifras contra el original.').italic = True
    elementos = [(p, y, 0, texto) for p, texto, y in contenido.texto]
    for tabla in contenido.tablas:
        pagina, bbox = tabla.zonas[0] if tabla.zonas else (tabla.paginas[0], (0, 0, 0, 0))
        elementos.append((pagina, bbox[1], 1, tabla))
    pagina_actual = None
    for pagina, _, tipo, elemento in sorted(elementos, key=lambda e: (e[0], e[1], e[2])):
        if pagina_actual is not None and pagina != pagina_actual:
            documento.add_page_break()
        pagina_actual = pagina
        if tipo == 0:
            documento.add_paragraph(elemento)
            continue
        filas = ([elemento.encabezado] if elemento.encabezado else []) + elemento.filas
        ancho = max(len(f) for f in filas)
        tabla_word = documento.add_table(rows=len(filas), cols=ancho)
        tabla_word.style = 'Table Grid'
        for r, fila in enumerate(filas):
            for c, valor in enumerate(fila):
                tabla_word.cell(r, c).text = valor
        if elemento.encabezado:
            for celda in tabla_word.rows[0].cells:
                for run in celda.paragraphs[0].runs:
                    run.bold = True
    buf = io.BytesIO()
    documento.save(buf)
    return buf.getvalue()


def pdf_a_docx(ruta_pdf):
    """
    PDF digital: pdf2docx (conserva diseño, tablas e imágenes).
    PDF escaneado (la mayoría de páginas son imagen): pdf2docx solo produciría
    imágenes, así que se reconstruye un Word editable con OCR.
    """
    with fitz.open(ruta_pdf) as doc:
        total = doc.page_count
        escaneadas = sum(1 for page in doc if es_escaneada(page))
    if total and escaneadas / total >= 0.5:
        if not ruta_tessdata():
            raise ValueError('El PDF es escaneado y el OCR no está instalado: no se puede extraer texto editable.')
        return _docx_desde_contenido(extraer_contenido(ruta_pdf))

    with tempfile.TemporaryDirectory() as tmp:
        ruta_out = os.path.join(tmp, 'convertido.docx')
        cv = Converter(ruta_pdf)
        try:
            cv.convert(ruta_out, start=0, end=None)
        finally:
            cv.close()
        with open(ruta_out, 'rb') as f:
            return f.read()


# ── Office → PDF ──────────────────────────────────────────────────

def ruta_libreoffice():
    """Ruta del ejecutable de LibreOffice o None si no está instalado."""
    configurada = os.environ.get('SOFFICE_PATH')
    if configurada and os.path.isfile(configurada):
        return configurada
    if sys.platform.startswith('win'):
        for base in (os.environ.get('PROGRAMFILES'), os.environ.get('PROGRAMFILES(X86)')):
            if not base:
                continue
            # soffice.com espera a que termine la conversión; soffice.exe retorna de inmediato.
            for exe in ('soffice.com', 'soffice.exe'):
                candidata = os.path.join(base, 'LibreOffice', 'program', exe)
                if os.path.isfile(candidata):
                    return candidata
    for nombre in ('soffice', 'libreoffice'):
        encontrada = shutil.which(nombre)
        if encontrada:
            return encontrada
    return None


def convertir_con_libreoffice(ruta_entrada, timeout=180):
    """Convierte un documento Office a PDF con LibreOffice headless. None si no está disponible o falla."""
    soffice = ruta_libreoffice()
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        salida = os.path.join(tmp, 'salida')
        os.makedirs(salida)
        # Perfil aislado: permite conversiones simultáneas y no toca el perfil del usuario.
        perfil = pathlib.Path(tmp, 'perfil').as_uri()
        cmd = [
            soffice, '--headless', '--norestore', '--nolockcheck', '--nodefault',
            '--nofirststartwizard', f'-env:UserInstallation={perfil}',
            '--convert-to', 'pdf', '--outdir', salida, ruta_entrada,
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=timeout, check=False)
        except (subprocess.TimeoutExpired, OSError) as e:
            log.warning('LibreOffice no completó la conversión: %s', e)
            return None
        pdfs = [f for f in os.listdir(salida) if f.lower().endswith('.pdf')]
        if not pdfs:
            log.warning('LibreOffice no generó PDF (código %s): %s', proc.returncode,
                        proc.stderr.decode(errors='replace')[:500])
            return None
        with open(os.path.join(salida, pdfs[0]), 'rb') as f:
            return f.read()


def _parrafo(texto, estilo):
    return Paragraph(escape(texto).replace('\n', '<br/>'), estilo)


def _docx_a_pdf_basico(ruta_docx):
    """Renderizado básico (sin LibreOffice): párrafos, títulos, tablas e imágenes."""
    doc = Document(ruta_docx)
    buf = io.BytesIO()
    pdf_doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=50, rightMargin=50,
                                topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    celda = styles['BodyText'].clone('celda', fontSize=8, leading=10)
    story = []
    ancho_util = A4[0] - 100

    paragraph_by_element = {para._p: para for para in doc.paragraphs}

    for element in doc.element.body.iterchildren():
        if element.tag.endswith('}tbl'):
            tabla = DocxTable(element, doc)
            filas = []
            for row in tabla.rows:
                filas.append([_parrafo(c.text.strip(), celda) for c in row.cells])
            if filas:
                n_cols = max(len(f) for f in filas)
                filas = [f + [''] * (n_cols - len(f)) for f in filas]
                t = Table(filas, colWidths=[ancho_util / n_cols] * n_cols, repeatRows=1)
                t.setStyle(TableStyle([
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#9AA5BE')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ]))
                story += [t, Spacer(1, 8)]
            continue

        if not element.tag.endswith('}p'):
            continue
        para = paragraph_by_element.get(element)
        if para is None:
            continue

        texto = para.text.strip()
        if texto:
            nombre_estilo = para.style.name if para.style is not None else ''
            if nombre_estilo.startswith('Heading 1') or nombre_estilo == 'Title':
                estilo = styles['Heading1']
            elif nombre_estilo.startswith('Heading'):
                estilo = styles['Heading2']
            else:
                estilo = styles['Normal']
            story += [_parrafo(texto, estilo), Spacer(1, 4)]

        # Imágenes inline del párrafo.
        for run in para.runs:
            for blip in run._element.xpath('.//a:blip'):
                rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                image_part = para.part.related_parts.get(rid) if rid else None
                if not image_part:
                    continue
                try:
                    img = RLImage(io.BytesIO(image_part.blob))
                except Exception:
                    continue
                if img.drawWidth > ancho_util:
                    escala = ancho_util / img.drawWidth
                    img.drawWidth = ancho_util
                    img.drawHeight *= escala
                story += [img, Spacer(1, 8)]

        if not texto and not para.runs:
            story.append(Spacer(1, 8))

    if not story:
        story = [Paragraph('(documento vacío)', styles['Normal'])]
    pdf_doc.build(story)
    return buf.getvalue()


def docx_a_pdf(ruta_docx):
    return convertir_con_libreoffice(ruta_docx) or _docx_a_pdf_basico(ruta_docx)


def _xlsx_a_pdf_basico(ruta_xlsx):
    """Renderizado básico (sin LibreOffice): todas las hojas visibles, valores calculados."""
    # data_only=True: muestra el valor calculado guardado por Excel, no la fórmula.
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True, read_only=True)
    styles = getSampleStyleSheet()
    celda = styles['BodyText'].clone('celda', fontSize=7, leading=8.5)
    story = []

    hojas = [ws for ws in wb.worksheets if getattr(ws, 'sheet_state', 'visible') == 'visible']
    max_cols_global = 1
    contenido = []
    for ws in hojas:
        filas = [['' if v is None else str(v) for v in fila] for fila in ws.iter_rows(values_only=True)]
        filas = [f for f in filas if any(c.strip() for c in f)]
        if filas:
            n_cols = max(len(f) for f in filas)
            # Quita columnas vacías al final.
            while n_cols > 1 and all((len(f) < n_cols or not f[n_cols - 1].strip()) for f in filas):
                n_cols -= 1
            filas = [(f + [''] * n_cols)[:n_cols] for f in filas]
            max_cols_global = max(max_cols_global, n_cols)
        contenido.append((ws.title, filas))
    wb.close()

    pagesize = landscape(A4) if max_cols_global > 8 else A4
    ancho_util = pagesize[0] - 60
    buf = io.BytesIO()
    pdf_doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=30, rightMargin=30,
                                topMargin=40, bottomMargin=40)

    for idx, (titulo, filas) in enumerate(contenido):
        if idx:
            story.append(PageBreak())
        story += [_parrafo(titulo, styles['Heading2']), Spacer(1, 6)]
        if not filas:
            story.append(Paragraph('(hoja vacía)', styles['Normal']))
            continue
        n_cols = len(filas[0])
        # Ancho proporcional al contenido más largo de cada columna (acotado).
        largos = [min(40, max(4, max(len(f[c]) for f in filas))) for c in range(n_cols)]
        total = sum(largos)
        anchos = [ancho_util * l / total for l in largos]
        data = [[_parrafo(v, celda) for v in f] for f in filas]
        tabla = Table(data, colWidths=anchos, repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B4FD8')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EBF2FF')]),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#C3D9FF')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(tabla)

    if not story:
        story = [Paragraph('(libro vacío)', styles['Normal'])]
    pdf_doc.build(story)
    return buf.getvalue()


def xlsx_a_pdf(ruta_xlsx):
    return convertir_con_libreoffice(ruta_xlsx) or _xlsx_a_pdf_basico(ruta_xlsx)


# ── Compresión ────────────────────────────────────────────────────

def _recomprimir_imagenes(doc, dpi_umbral, dpi_objetivo, calidad):
    """
    Reduce a dpi_objetivo (JPEG) las imágenes que se muestran por encima de dpi_umbral.
    Se omiten imágenes con transparencia, máscaras y bitonales (JBIG2/CCITT ya son
    eficientes y en JPEG crecerían). Retorna cuántas imágenes se reemplazaron.
    """
    procesadas = set()
    reemplazadas = 0
    for page in doc:
        for info in page.get_image_info(xrefs=True):
            xref = info.get('xref')
            bbox = fitz.Rect(info['bbox'])
            if not xref or xref in procesadas or bbox.is_empty or bbox.width < 1 or bbox.height < 1:
                continue
            procesadas.add(xref)
            dpi = max(info['width'] / (bbox.width / 72), info['height'] / (bbox.height / 72))
            if dpi <= dpi_umbral:
                continue
            try:
                if doc.xref_get_key(xref, 'SMask')[0] != 'null' or doc.xref_get_key(xref, 'Mask')[0] != 'null':
                    continue
                if doc.xref_get_key(xref, 'ImageMask')[1] == 'true' or doc.xref_get_key(xref, 'BitsPerComponent')[1] == '1':
                    continue
                pix = fitz.Pixmap(doc, xref)
                if pix.alpha:
                    continue
                if pix.n not in (1, 3):
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                factor = dpi_objetivo / dpi
                pix = fitz.Pixmap(pix, max(1, int(pix.width * factor)), max(1, int(pix.height * factor)), None)
                datos = pix.tobytes('jpeg', jpg_quality=calidad)
                if len(datos) >= len(doc.xref_stream_raw(xref) or b''):
                    continue
                page.replace_image(xref, stream=datos)
                reemplazadas += 1
            except Exception as e:
                log.debug('No se recomprimió la imagen xref=%s: %s', xref, e)
    return reemplazadas


def _rasterizar_paginas(doc, zoom, calidad, gris):
    """Convierte cada página en una imagen JPEG del mismo tamaño de página."""
    with fitz.open() as nuevo:
        espacio = fitz.csGRAY if gris else fitz.csRGB
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), colorspace=espacio, alpha=False)
            pagina = nuevo.new_page(width=page.rect.width, height=page.rect.height)
            pagina.insert_image(pagina.rect, stream=pix.tobytes('jpeg', jpg_quality=calidad))
        return nuevo.tobytes(garbage=4, deflate=True)


def comprimir_pdf(ruta_pdf, nivel='media', contrasena=''):
    """
    Comprime un PDF.
      baja:  limpieza sin pérdida (objetos huérfanos + deflate).
      media: sin pérdida + imágenes de más de 300 dpi a 200 dpi (JPEG 85).
      alta:  sin pérdida + imágenes de más de 170 dpi a 150 dpi (JPEG 70).
      ultra: cada página como imagen JPEG en escala de grises a ~94 dpi
             (máxima reducción; se pierde el texto seleccionable).
    Si el resultado no es menor que el original, se devuelve el original.
    """
    if nivel not in NIVELES_COMPRESION:
        raise ValueError(f'Nivel inválido. Opciones: {", ".join(NIVELES_COMPRESION)}')

    tamano_original = os.path.getsize(ruta_pdf)
    with abrir_pdf(ruta_pdf, contrasena) as doc:
        if nivel == 'ultra':
            resultado = _rasterizar_paginas(doc, zoom=1.3, calidad=72, gris=True)
        else:
            if nivel == 'media':
                _recomprimir_imagenes(doc, dpi_umbral=300, dpi_objetivo=200, calidad=85)
            elif nivel == 'alta':
                _recomprimir_imagenes(doc, dpi_umbral=170, dpi_objetivo=150, calidad=70)
            if nivel == 'baja':
                opts = {'garbage': 1, 'deflate': True}
            else:
                opts = {'garbage': 4, 'deflate': True, 'deflate_images': True,
                        'deflate_fonts': True, 'clean': True}
            resultado = doc.tobytes(**opts)

    if len(resultado) >= tamano_original:
        log.info('Compresión %s sin mejora (%d -> %d bytes); se conserva el original',
                 nivel, tamano_original, len(resultado))
        with open(ruta_pdf, 'rb') as f:
            return f.read()

    log.info('Compresión %s: %.1f KB -> %.1f KB (%.1f %%)', nivel, tamano_original / 1024,
             len(resultado) / 1024, (1 - len(resultado) / tamano_original) * 100)
    return resultado


# ── Cifrado ───────────────────────────────────────────────────────

def cifrar_pdf(ruta_pdf, contrasena):
    """Cifra con AES-256 (el cifrado RC4 de 128 bits anterior está obsoleto)."""
    with abrir_pdf(ruta_pdf) as doc:
        return _guardar(
            doc,
            encryption=fitz.PDF_ENCRYPT_AES_256,
            user_pw=contrasena,
            owner_pw=contrasena,
        )


# ── Aplanado ──────────────────────────────────────────────────────

def aplanar_pdf(ruta_pdf):
    """Fija anotaciones (texto, firmas, dibujos) y campos de formulario como contenido de la página."""
    with abrir_pdf(ruta_pdf) as doc:
        doc.bake(annots=True, widgets=True)
        return _guardar(doc)
