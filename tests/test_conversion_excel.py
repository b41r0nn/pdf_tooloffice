import io
from datetime import datetime

import openpyxl
import pymupdf as fitz
import pytest

from app import conversion_excel as ce
from app.conversion_excel import (
    Contenido,
    Tabla,
    _es_encabezado,
    _orden_fechas,
    _votar_decimal,
    a_fecha,
    a_numero,
    escribir_excel,
    extraer_contenido,
    limpiar,
    pdf_a_xlsx,
    tipar_columna,
)

# ── Números y fechas ──────────────────────────────────────────────

@pytest.mark.parametrize('texto, decimal, esperado', [
    ('1.250.000,00', ',', (1250000.0, '#,##0.00')),
    ('1,250,000.00', '.', (1250000.0, '#,##0.00')),
    ('$ 3.016.828,92', ',', (3016828.92, '#,##0.00')),
    ('(1.234,50)', ',', (-1234.5, '#,##0.00')),
    ('1.234,00-', ',', (-1234.0, '#,##0.00')),
    ('19%', ',', (0.19, '0%')),
    ('450,5', ',', (450.5, '#,##0.0')),
    ('12', ',', (12, 'General')),
    ('4.775', ',', (4775, '#,##0')),
    ('1.25', '.', (1.25, '#,##0.00')),
])
def test_a_numero(texto, decimal, esperado):
    valor, formato = a_numero(texto, decimal)
    assert valor == pytest.approx(esperado[0]) and formato == esperado[1]


@pytest.mark.parametrize('texto', ['000123', '900.123.456-7', '1234567890123456', 'abc', '1.25', 'RF-1000', ''])
def test_no_son_numeros(texto):
    # '1.25' con coma decimal es ambiguo (grupo de miles incompleto): queda como texto
    assert a_numero(texto, ',') is None


def test_votar_decimal():
    assert _votar_decimal(['18.500,00', '450,50', 'Guantes']) == ','
    assert _votar_decimal(['6,612.02', '2,618,359.92']) == '.'
    assert _votar_decimal(['4,775', '396']) is None


@pytest.mark.parametrize('texto, esperado', [
    ('15/10/2026', datetime(2026, 10, 15)),
    ('2026-10-15', datetime(2026, 10, 15)),
    ('05-ago-2026', datetime(2026, 8, 5)),
    ('15 de octubre de 2026', datetime(2026, 10, 15)),
    ('15/10/26', datetime(2026, 10, 15)),
    ('15/10/2026 14:30', datetime(2026, 10, 15, 14, 30)),
])
def test_a_fecha(texto, esperado):
    assert a_fecha(texto) == esperado


def test_fechas_invalidas_y_orden():
    assert a_fecha('31/02/2026') is None
    assert a_fecha('1.250') is None
    assert _orden_fechas(['10/15/2026']) is False
    assert a_fecha('10/15/2026', dia_primero=False) == datetime(2026, 10, 15)


def test_tipar_columna_codigos_quedan_texto():
    assert tipar_columna(['000123', '004567'], 'Código') == [('000123', None), ('004567', None)]
    assert tipar_columna(['455589', '120429'], 'Referencia') == [('455589', None), ('120429', None)]


def test_tipar_columna_formato_unico_y_umbral():
    col = tipar_columna(['396', '4,775', '1,060'], 'Existencia', '.')
    assert [v for v, _ in col] == [396, 4775, 1060]
    assert {f for _, f in col} == {'#,##0'}
    mezclada = tipar_columna(['1.000', 'N/A', 'sin dato', '2.000'], 'Valor', ',')
    assert all(isinstance(v, str) for v, _ in mezclada)   # menos del 80 % numérico


def test_acumulado_no_es_columna_de_codigo():
    col = tipar_columna(['1.000,00', '2.000,00'], 'Valor acumulado', ',')
    assert [v for v, _ in col] == [1000.0, 2000.0]


def test_es_encabezado():
    assert _es_encabezado(['Concepto', '2025', '2024'])
    assert not _es_encabezado(['RF-1000', 'Catéter', 'L52445', '05/07/2027'])
    assert not _es_encabezado(['', '', 'Total'])


def test_limpiar():
    assert limpiar('Guantes de\nnitrilo') == 'Guantes de nitrilo'
    assert limpiar('antisép-\ntico') == 'antiséptico'
    assert limpiar('19% |') == '19%'
    assert limpiar(None) == ''


# ── PDFs de prueba ────────────────────────────────────────────────

def _tabla_con_bordes(page, x0, y0, filas, anchos, alto=18):
    y = y0
    for fila in filas:
        x = x0
        for texto, ancho in zip(fila, anchos):
            page.draw_rect(fitz.Rect(x, y, x + ancho, y + alto), color=(0, 0, 0), width=0.6)
            page.insert_text((x + 3, y + 13), texto, fontsize=9)
            x += ancho
        y += alto


ENCABEZADO = ['Código', 'Producto', 'Cant.', 'Valor']
ANCHOS = [70, 180, 50, 90]


def _pdf_factura(paginas=1, filas_por_pagina=4):
    with fitz.open() as doc:
        n = 1
        for p in range(paginas):
            page = doc.new_page()
            if p == 0:
                page.insert_text((72, 60), 'Factura No. FE-1', fontsize=12)
                page.insert_text((72, 80), 'NIT: 900.123.456-7', fontsize=10)
                page.insert_text((300, 80), 'Fecha: 15/10/2026', fontsize=10)
            filas = [ENCABEZADO]
            for _ in range(filas_por_pagina):
                filas.append([f'{n:06d}', f'Producto {n}', str(n * 2), f'{n * 1000:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')])
                n += 1
            _tabla_con_bordes(page, 72, 120, filas, ANCHOS)
            if p == paginas - 1:
                page.insert_text((72, 120 + 18 * (len(filas) + 1)), 'TOTAL A PAGAR: $ 1.234.567,89', fontsize=10)
        return doc.tobytes()


def _libro(datos):
    return openpyxl.load_workbook(io.BytesIO(datos))


def test_factura_con_bordes_a_excel(tmp_path):
    ruta = tmp_path / 'f.pdf'
    ruta.write_bytes(_pdf_factura())
    wb = _libro(pdf_a_xlsx(str(ruta), 'f.pdf'))
    assert wb.sheetnames == ['Resumen', 'Tabla 1', 'Datos', 'Texto']
    ws = wb['Tabla 1']
    assert [c.value for c in ws[1]] == ENCABEZADO
    assert ws['A2'].value == '000001' and ws['A2'].data_type == 's'      # código conserva ceros
    assert ws['C2'].value == 2 and ws['D3'].value == pytest.approx(2000.0)
    assert ws['D2'].number_format == '#,##0.00'
    assert 'Tabla1' in ws.tables
    datos = {fila[1]: fila[2] for fila in wb['Datos'].iter_rows(min_row=2, values_only=True)}
    assert datos['NIT'] == '900.123.456-7'
    assert datos['Fecha'] == datetime(2026, 10, 15)
    assert datos['TOTAL A PAGAR'] == pytest.approx(1234567.89)


def test_tabla_en_varias_paginas_se_une_sin_encabezados_repetidos(tmp_path):
    ruta = tmp_path / 'multi.pdf'
    ruta.write_bytes(_pdf_factura(paginas=3, filas_por_pagina=5))
    contenido = extraer_contenido(str(ruta))
    assert len(contenido.tablas) == 1
    tabla = contenido.tablas[0]
    assert tabla.paginas == [1, 2, 3] and len(tabla.filas) == 15
    assert ENCABEZADO not in tabla.filas
    # El contenido de la tabla en las páginas 2 y 3 no se duplica en Texto
    assert not any('Producto 7' in t for _, t, _ in contenido.texto)


def test_sin_tablas_deja_texto_y_aviso(tmp_path):
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), 'Memorando interno', fontsize=12)
        page.insert_text((72, 100), 'Responsable: Área de compras', fontsize=10)
        ruta = tmp_path / 'm.pdf'
        doc.save(ruta)
    wb = _libro(pdf_a_xlsx(str(ruta)))
    assert 'Tabla 1' not in wb.sheetnames and 'Texto' in wb.sheetnames
    resumen = ' '.join(str(c.value) for fila in wb['Resumen'].iter_rows() for c in fila if c.value)
    assert 'No se detectaron tablas' in resumen


def test_excel_seguro_con_caracteres_invalidos_y_formulas():
    contenido = Contenido(paginas=1, tablas=[Tabla(paginas=[1], encabezado=['A', 'A'], filas=[
        ['=HYPERLINK("http://x")', 'texto\x07con control'], ['x' * 40000, '1']])],
        texto=[(1, '=1+1', 0)], datos=[(1, 'Campo', '=cmd')])
    wb = _libro(escribir_excel(contenido))
    ws = wb['Tabla 1']
    assert [c.value for c in ws[1]] == ['A', 'A (2)']                      # encabezados únicos
    assert ws['A2'].data_type == 's' and ws['A2'].value.startswith('=')
    assert ws['B2'].value == 'textocon control'
    assert len(ws['A3'].value) == 32767
    assert wb['Texto']['B2'].data_type == 's'
    assert wb['Datos']['C2'].data_type == 's'


@pytest.mark.skipif(ce.camelot is None, reason='Camelot no instalado')
def test_tabla_sin_bordes_con_camelot(tmp_path):
    with fitz.open() as doc:
        page = doc.new_page()
        columnas = [72, 160, 330, 420]
        filas = [['Fecha', 'Descripción', 'Débito', 'Saldo']] + [
            [f'{d:02d}/09/2026', f'Movimiento {d}', f'{d * 10},50', f'1.{d:03d}.000,00'] for d in range(1, 16)]
        for i, fila in enumerate(filas):
            for x, texto in zip(columnas, fila):
                page.insert_text((x, 100 + i * 16), texto, fontsize=9)
        ruta = tmp_path / 'extracto.pdf'
        doc.save(ruta)
    contenido = extraer_contenido(str(ruta))
    assert len(contenido.tablas) == 1
    assert contenido.tablas[0].encabezado == ['Fecha', 'Descripción', 'Débito', 'Saldo']
    assert len(contenido.tablas[0].filas) == 15
    ws = _libro(escribir_excel(contenido))['Tabla 1']
    assert ws['A2'].value == datetime(2026, 9, 1)
    assert ws['C2'].value == pytest.approx(10.5)


@pytest.mark.skipif(ce.ruta_tessdata() is None, reason='OCR (tessdata spa) no instalado')
def test_pdf_escaneado_usa_ocr(tmp_path):
    with fitz.open('pdf', _pdf_factura()) as original, fitz.open() as escaneado:
        pix = original[0].get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        page = escaneado.new_page(width=original[0].rect.width, height=original[0].rect.height)
        page.insert_image(page.rect, stream=pix.tobytes('png'))
        ruta = tmp_path / 'escaneado.pdf'
        escaneado.save(ruta)
    contenido = extraer_contenido(str(ruta))
    assert contenido.paginas_ocr == [1]
    assert any('NIT' == campo for _, campo, _ in contenido.datos)
    assert contenido.tablas and any('000002' in fila for fila in contenido.tablas[0].filas)


@pytest.mark.skipif(ce.ruta_tessdata() is None, reason='OCR (tessdata spa) no instalado')
def test_pdf_escaneado_a_word_editable(tmp_path):
    from docx import Document
    from app.services import pdf_a_docx
    with fitz.open('pdf', _pdf_factura()) as original, fitz.open() as escaneado:
        pix = original[0].get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        page = escaneado.new_page(width=original[0].rect.width, height=original[0].rect.height)
        page.insert_image(page.rect, stream=pix.tobytes('png'))
        ruta = tmp_path / 'escaneado.pdf'
        escaneado.save(ruta)
    documento = Document(io.BytesIO(pdf_a_docx(str(ruta))))
    texto = '\n'.join(p.text for p in documento.paragraphs)
    assert 'NIT' in texto and len(documento.tables) == 1


def test_datos_ignora_parrafos_y_separa_etiquetas(tmp_path):
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), 'PRIMERA. OBJETO: El proveedor se obliga a suministrar insumos médicos de acuerdo', fontsize=9)
        page.insert_text((72, 100), 'Firma: ____________   Fecha: 01/10/2026', fontsize=10)
        page.insert_text((72, 120), 'Titular: REDIHOS S.A.S. · Periodo: septiembre 2026', fontsize=10)
        ruta = tmp_path / 'c.pdf'
        doc.save(ruta)
    datos = {campo: valor for _, campo, valor in extraer_contenido(str(ruta)).datos}
    assert datos == {'Fecha': '01/10/2026', 'Titular': 'REDIHOS S.A.S.', 'Periodo': 'septiembre 2026'}
